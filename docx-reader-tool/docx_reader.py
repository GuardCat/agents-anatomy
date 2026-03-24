#!/usr/bin/env -S uv run
# /// script
# dependencies = [
#     "python-docx",
# ]
# ///

import sys
import os
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def parse_docx(file_path, output_dir="extracted_images"):
    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"Ошибка чтения файла: {e}")
        return

    # Создаем директорию для картинок, если её нет
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    print(f"# Анализ документа: {file_path}\n")

    print("## Текст")
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)

    print("\n## Ссылки")
    rels = doc.part.rels
    for rel in rels.values():
        if rel.reltype == RT.HYPERLINK:
            print(f"- {rel.target_ref}")

    print("\n## Изображения")
    for idx, shape in enumerate(doc.inline_shapes):
        w = shape.width.cm if shape.width else 0
        h = shape.height.cm if shape.height else 0
        
        try:
            # Извлекаем ID связи (rId) для текущей картинки
            blip = shape._inline.graphic.graphicData.pic.blipFill.blip
            rId = blip.embed
            
            # Получаем сам файл картинки из внутренностей docx
            image_part = doc.part.related_parts[rId]
            
            # Определяем формат (jpeg, png и т.д.)
            ext = image_part.content_type.split('/')[-1]
            if ext == 'jpeg': ext = 'jpg'
            
            # Формируем путь и сохраняем
            filename = f"image_{idx + 1}.{ext}"
            filepath = os.path.join(output_dir, filename)
            
            with open(filepath, "wb") as f:
                f.write(image_part.blob)
                
            print(f"- Картинка {idx + 1}: {w:.2f} см x {h:.2f} см. Сохранена в: `{filepath}`")
        except Exception as e:
            print(f"- Картинка {idx + 1}: {w:.2f} см x {h:.2f} см. Ошибка извлечения файла: {e}")

    print("\n## Таблицы")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n### Таблица {t_idx + 1}")
        for r_idx, row in enumerate(table.rows):
            row_text = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
            print("| " + " | ".join(row_text) + " |")
            if r_idx == 0:
                print("|" + "|".join(["---"] * len(row.cells)) + "|")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        
        # По умолчанию создаем папку с именем документа + _images
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        out_dir = f"{base_name}_images"
        
        # Если передан второй аргумент, используем его
        if len(sys.argv) > 2:
            out_dir = sys.argv[2]
            
        parse_docx(file_path, out_dir)
    else:
        print("Использование: uv run docx_reader.py <путь_к_docx> [путь_к_папке_для_картинок]")
